from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from carteis.domain.errors import ValidationError
from carteis.services import ClientService, DriverService, SpreadsheetService
from carteis.services.spreadsheets import CLIENT_HEADERS, DRIVER_HEADERS
from tests.factories import make_client, make_driver


@pytest.fixture
def sheets(conn):
    return SpreadsheetService(conn)


def _write(path, rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    return path


def test_export_drivers(conn, sheets, tmp_path):
    DriverService(conn).register(make_driver())
    DriverService(conn).register(make_driver(dni="00000000T", surname="Abad"))
    DriverService(conn).deactivate(2, on=date(2024, 1, 31))

    path = tmp_path / "drivers.xlsx"
    assert sheets.export_drivers(path) == 2

    sheet = load_workbook(path).active
    assert [cell.value for cell in sheet[1]] == DRIVER_HEADERS
    abad = [cell.value for cell in sheet[2]]
    assert abad[1] == "00000000T"
    assert abad[2] == datetime(2021, 9, 15)  # real Excel dates, not text
    assert abad[8] == "612345678"  # phone stays text
    assert abad[9] == 1500.5  # salary is a number
    assert abad[10] == "A-B"
    assert abad[11] == datetime(2024, 1, 31)
    assert sheet["C2"].number_format == "DD/MM/YYYY"
    assert sheet["J2"].number_format == '#,##0.00 "€"'
    assert sheet.freeze_panes == "A2"


def test_export_clients(conn, sheets, tmp_path):
    ClientService(conn).register(make_client())
    path = tmp_path / "clients.xlsx"
    assert sheets.export_clients(path) == 1
    rows = list(load_workbook(path).active.iter_rows(values_only=True))
    assert list(rows[0]) == CLIENT_HEADERS
    assert rows[1][1:] == (
        "53854586V",
        "So Factive",
        "Estrela 1",
        "619669207",
        "Pontevedra",
        "Cangas",
        None,
    )


def test_export_then_import_round_trip(conn, sheets, tmp_path):
    drivers = DriverService(conn)
    drivers.register(make_driver())
    drivers.register(make_driver(dni="00000000T", surname="Abad"))
    drivers.deactivate(2, on=date(2024, 1, 31))
    ClientService(conn).register(make_client())
    sheets.export_drivers(tmp_path / "d.xlsx")
    sheets.export_clients(tmp_path / "c.xlsx")
    before = drivers.list(), ClientService(conn).list()

    conn.execute("DELETE FROM drivers")
    conn.execute("DELETE FROM clients")
    assert sheets.import_drivers(tmp_path / "d.xlsx").inserted == 2
    assert sheets.import_clients(tmp_path / "c.xlsx").inserted == 1

    after = drivers.list(), ClientService(conn).list()
    strip_ids = lambda people: [(p.dni, p.left_on, getattr(p, "salary", None)) for p in people]  # noqa: E731
    assert strip_ids(after[0]) == strip_ids(before[0])
    assert strip_ids(after[1]) == strip_ids(before[1])


def test_import_reports_each_rejected_row(conn, sheets, tmp_path):
    DriverService(conn).register(make_driver())  # 12345678Z already exists
    path = _write(
        tmp_path / "in.xlsx",
        [
            # Different order, accents/case variations and no ID or dates: still recognised.
            [
                "dni",
                "APELLIDOS",
                "Nombre",
                "Direccion",
                "Provincia",
                "Localidad",
                "movil",
                "Salario",
                "Licencias",
            ],
            [
                "00000000T",
                "Abad",
                "Luis",
                "Rúa 1",
                "Pontevedra",
                "Vigo",
                612345678.0,
                "1.500,50",
                "A-B",
            ],
            [
                "X1234567L",
                "Ruiz",
                "Eva",
                "Rúa 2",
                "Pontevedra",
                "Cangas",
                "612345679",
                35000.0,
                "c",
            ],
            [None, None, None, None, None, None, None, None, None],  # blank row: ignored
            ["12345678Z", "Dup", "Dup", "x", "Pontevedra", "Vigo", "612345678", 1, "A"],
            ["12345678A", "Bad", "Dni", "x", "Pontevedra", "Vigo", "612345678", 1, "A"],
            ["Y1234567X", "Bad", "Place", "x", "A Coruña", "Vigo", "612345678", 1, "A"],
            ["Z1234567R", "Bad", "Salary", "x", "Pontevedra", "Vigo", "612345678", "-5", "A"],
        ],
    )
    result = sheets.import_drivers(path)

    assert result.inserted == 2
    assert result.total == 6
    assert [(r.row, r.reason) for r in result.rejected] == [
        # Header is row 1 and row 4 is blank, so these match what Excel shows.
        (5, "Ya existe un conductor con ese DNI."),
        (6, "El DNI/NIE no es válido."),
        (7, "La localidad «Vigo» no existe en la provincia «A Coruña»."),
        (8, "Importe no válido: usa solo números y como máximo dos decimales."),
    ]
    abad = DriverService(conn).find_by_dni("00000000T")
    assert (abad.phone, abad.salary, abad.hired_on) == ("612345678", Decimal("1500.50"), None)
    assert {str(item) for item in DriverService(conn).find_by_dni("X1234567L").licenses} == {"C"}


def test_import_parses_text_and_real_dates(conn, sheets, tmp_path):
    path = _write(
        tmp_path / "in.xlsx",
        [
            DRIVER_HEADERS,
            [
                None,
                "00000000T",
                "15/09/2021",
                "A",
                "B",
                "C",
                "Pontevedra",
                "Vigo",
                "612345678",
                100,
                "A",
                datetime(2024, 1, 31),
            ],
            [
                None,
                "X1234567L",
                "2021-13-01",
                "A",
                "B",
                "C",
                "Pontevedra",
                "Vigo",
                "612345678",
                100,
                "A",
                None,
            ],
        ],
    )
    result = sheets.import_drivers(path)
    driver = DriverService(conn).find_by_dni("00000000T")
    assert (driver.hired_on, driver.left_on) == (date(2021, 9, 15), date(2024, 1, 31))
    assert result.rejected[0].reason.startswith("Fecha no válida")


def test_import_requires_columns(sheets, tmp_path):
    path = _write(tmp_path / "in.xlsx", [["DNI", "Nombre"]])
    with pytest.raises(ValidationError, match="Faltan columnas en el fichero: Apellidos"):
        sheets.import_drivers(path)


def test_import_rejects_non_xlsx(sheets, tmp_path):
    path = tmp_path / "old.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0 legacy binary")
    with pytest.raises(ValidationError, match="no es una hoja de cálculo .xlsx"):
        sheets.import_clients(path)


def test_import_clients_rejections(conn, sheets, tmp_path):
    path = _write(
        tmp_path / "in.xlsx",
        [
            ["DNI", "Razón social", "Dirección", "Teléfono", "Provincia", "Localidad"],
            ["53854586V", "So Factive", "Estrela", 619669207, "Pontevedra", "Vigo"],
            ["53857823B", "", "Pza 40", 688957121, "Pontevedra", "Vigo"],
        ],
    )
    result = sheets.import_clients(path)
    assert result.inserted == 1
    assert result.rejected[0].reason == "El campo «Razón social» es obligatorio."
    assert ClientService(conn).find_by_dni("53854586V").phone == "619669207"
