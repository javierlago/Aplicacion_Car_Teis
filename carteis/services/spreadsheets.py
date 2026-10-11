"""Excel (.xlsx) import and export of drivers and clients.

Columns are matched by header, not by position, so an exported file can be edited and
imported back. Headers are compared ignoring case, accents and surrounding spaces.
"""

import sqlite3
import unicodedata
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from carteis.data.database import transaction
from carteis.data.repositories import ClientRepository, DriverRepository
from carteis.domain.errors import DomainError, ValidationError
from carteis.domain.formatting import parse_date
from carteis.domain.models import Client, Driver
from carteis.domain.validators import parse_money

_DATE_FORMAT = "DD/MM/YYYY"
_MONEY_FORMAT = '#,##0.00 "€"'

DRIVER_HEADERS = [
    "ID",
    "DNI",
    "Fecha alta",
    "Apellidos",
    "Nombre",
    "Dirección",
    "Provincia",
    "Localidad",
    "Móvil",
    "Salario",
    "Licencias",
    "Fecha baja",
]
CLIENT_HEADERS = [
    "ID",
    "DNI",
    "Razón social",
    "Dirección",
    "Teléfono",
    "Provincia",
    "Localidad",
    "Fecha baja",
]
# Columns that may be missing on import ("ID" is always ignored: the database assigns it).
_OPTIONAL_HEADERS = {"ID", "Fecha alta", "Fecha baja"}


@dataclass
class RejectedRow:
    row: int  # spreadsheet row number, as Excel shows it
    reason: str


@dataclass
class ImportResult:
    inserted: int = 0
    rejected: list[RejectedRow] = field(default_factory=list)

    @property
    def total(self) -> int:
        return self.inserted + len(self.rejected)


class SpreadsheetService:
    def __init__(self, conn: sqlite3.Connection) -> None:
        self._conn = conn
        self._drivers = DriverRepository(conn)
        self._clients = ClientRepository(conn)

    def export_drivers(self, path: Path) -> int:
        rows = [
            [
                driver.id,
                driver.dni,
                driver.hired_on,
                driver.surname,
                driver.name,
                driver.address,
                driver.province,
                driver.municipality,
                driver.phone,
                driver.salary,
                "-".join(sorted(driver.licenses)),
                driver.left_on,
            ]
            for driver in self._drivers.list()
        ]
        _write_sheet(path, "Conductores", DRIVER_HEADERS, rows, money_columns={"Salario"})
        return len(rows)

    def export_clients(self, path: Path) -> int:
        rows = [
            [
                client.id,
                client.dni,
                client.company_name,
                client.address,
                client.phone,
                client.province,
                client.municipality,
                client.left_on,
            ]
            for client in self._clients.list()
        ]
        _write_sheet(path, "Clientes", CLIENT_HEADERS, rows)
        return len(rows)

    def import_drivers(self, path: Path) -> ImportResult:
        return self._import(path, DRIVER_HEADERS, _row_to_driver, self._drivers.add)

    def import_clients(self, path: Path) -> ImportResult:
        return self._import(path, CLIENT_HEADERS, _row_to_client, self._clients.add)

    def _import(
        self,
        path: Path,
        headers: list[str],
        build: Callable[[dict[str, Any]], Any],
        save: Callable[[Any], Any],
    ) -> ImportResult:
        """Insert every valid row; invalid or duplicate rows are reported, not inserted."""
        result = ImportResult()
        with transaction(self._conn):
            for number, values in _read_sheet(path, headers):
                try:
                    save(build(values))
                    result.inserted += 1
                except DomainError as error:
                    result.rejected.append(RejectedRow(number, str(error)))
        return result


def _write_sheet(
    path: Path,
    title: str,
    headers: list[str],
    rows: Iterable[list[Any]],
    money_columns: set[str] = frozenset(),
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for column, header in enumerate(headers, start=1):
        letter = sheet.cell(row=1, column=column).column_letter
        for (cell,) in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
            if isinstance(cell.value, date):
                cell.number_format = _DATE_FORMAT
            elif header in money_columns:
                cell.number_format = _MONEY_FORMAT
        sheet.column_dimensions[letter].width = max(12, len(header) + 4)
    sheet.freeze_panes = "A2"
    workbook.save(path)


def _normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).casefold().strip()
    return "".join(char for char in text if not unicodedata.combining(char))


def _read_sheet(path: Path, headers: list[str]) -> Iterable[tuple[int, dict[str, Any]]]:
    """Yield ``(row number, {header: value})`` for each non-empty data row."""
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises several unrelated types for bad files
        raise ValidationError("El fichero no es una hoja de cálculo .xlsx válida.") from error
    try:
        rows = workbook.active.iter_rows(values_only=True)
        header_row = next(rows, ())
        wanted = {_normalize_header(header): header for header in headers}
        positions = {
            wanted[name]: index
            for index, value in enumerate(header_row)
            if (name := _normalize_header(value)) in wanted
        }
        missing = [h for h in headers if h not in positions and h not in _OPTIONAL_HEADERS]
        if missing:
            raise ValidationError(f"Faltan columnas en el fichero: {', '.join(missing)}.")

        for number, row in enumerate(rows, start=2):
            if all(value in (None, "") for value in row):
                continue
            yield (
                number,
                {
                    header: row[index] if index < len(row) else None
                    for header, index in positions.items()
                },
            )
    finally:
        workbook.close()


def _text(value: Any) -> str:
    # openpyxl reads whole numbers as int (phones typed as numbers become "612345678").
    return "" if value is None else str(value).strip()


def _optional_date(value: Any, field_name: str) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):  # openpyxl returns datetime for every date cell
        return value.date()
    return parse_date(str(value), field=field_name)


def _money(value: Any) -> Decimal:
    if isinstance(value, int | float | Decimal):
        return parse_money(str(Decimal(str(value))), field="salary")
    return parse_money(_text(value), field="salary")


def _licenses(value: Any) -> frozenset[str]:
    return frozenset(char for char in _text(value).upper() if char.isalpha())


def _row_to_driver(values: dict[str, Any]) -> Driver:
    return Driver(
        dni=_text(values["DNI"]),
        hired_on=_optional_date(values.get("Fecha alta"), "hired_on"),
        surname=_text(values["Apellidos"]),
        name=_text(values["Nombre"]),
        address=_text(values["Dirección"]),
        province=_text(values["Provincia"]),
        municipality=_text(values["Localidad"]),
        phone=_text(values["Móvil"]),
        salary=_money(values["Salario"]),
        licenses=_licenses(values["Licencias"]),
        left_on=_optional_date(values.get("Fecha baja"), "left_on"),
    )


def _row_to_client(values: dict[str, Any]) -> Client:
    return Client(
        dni=_text(values["DNI"]),
        company_name=_text(values["Razón social"]),
        address=_text(values["Dirección"]),
        phone=_text(values["Teléfono"]),
        province=_text(values["Provincia"]),
        municipality=_text(values["Localidad"]),
        left_on=_optional_date(values.get("Fecha baja"), "left_on"),
    )
